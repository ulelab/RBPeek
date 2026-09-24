#!/usr/bin/env python3
"""
Build an inference BED for intersect_inference_bed.py from proximity-CLIP (APEX) peak calls.

In proximity CLIP, peak cDNA reflects the proportion of a transcript in the labelled
compartment rather than a discrete binding site, so the inference loci are restricted to
genes enriched in the experimental samples relative to the control:

  1. Gene bodies are taken from the annotation ('gene' rows of the GTF); chromosome names are
     converted to UCSC style (chr-prefixed; MT -> chrM) and non-primary contigs are removed.
  2. For each sample, the cDNA of every peak on the same strand as a gene is assigned to that
     gene in proportion to the fraction of the peak's width inside it, and gene totals are
     expressed per million of the sample's peak cDNA (mitochondrial peaks excluded).
  3. Per-gene means are taken over the experimental and over the control samples, and
     ratio = (mean_experimental + pseudocount) / (mean_control + pseudocount).
  4. A gene is retained if it appears in the differential-expression list (log2FoldChange
     > --min-lfc and pvalue <= --max-p) and ratio >= --min-ratio.
  5. Peaks from the experimental samples are pooled, overlapping peaks on the same strand are
     merged, each merged region is reduced to its 1 nt midpoint, and anchors within a retained
     gene on the same strand are written.

Output BED6:
    chrom  start  end (= start + 1)  name  score  strand
where name is <experimental>_<n>reps_<i> and score is the number of supporting replicates.

A per-gene table (<experimental>_gene_enrichment.tsv) lists every gene with a peak in either
group, with per-sample values, group means, the ratio and the retention decision.
"""

import argparse
import datetime
import glob
import os
import shutil
import subprocess
import sys
import tarfile
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
PRIMARY = {"chr%s" % c for c in list(range(1, 23)) + ["X", "Y"]}
CHRM = {"chrM", "chrMT", "MT", "M"}

# Gene symbols that spreadsheets convert to dates, with their current HGNC names.
DATE_SYMBOLS = {
    (3, d): ["MARCH%d" % d, "MARCHF%d" % d] for d in range(1, 12)
}
DATE_SYMBOLS.update({(9, d): ["SEPT%d" % d, "SEPTIN%d" % d] for d in range(1, 16)})
DATE_SYMBOLS[(12, 1)] = ["DEC1", "DELEC1"]


def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--raw-dir", type=Path, default=REPO / "Centrosome" / "raw",
                   help="Directory of <sample>_genome.*_Peaks.bed files, or one containing a single "
                        ".tar.gz of them (default Centrosome/raw)")
    p.add_argument("--deseq", type=Path, required=True,
                   help="Differential-expression table (.xlsx; first sheet; columns: gene symbol, "
                        "baseMean, log2FoldChange, lfcSE, stat, pvalue, padj)")
    p.add_argument("-g", "--gtf", type=Path, required=True,
                   help="Gene annotation GTF of the assembly the peaks were called on (Ensembl 109 for "
                        "the Centrosome analysis)")
    p.add_argument("--experimental", default="Centro_apex", help="Sample-name prefix of the experimental group")
    p.add_argument("--control", default="NT_apex", help="Sample-name prefix of the control group")
    p.add_argument("--min-lfc", type=float, default=1.0, help="Minimum log2FoldChange in --deseq (default 1)")
    p.add_argument("--max-p", type=float, default=0.05, help="Maximum pvalue in --deseq (default 0.05)")
    p.add_argument("--min-ratio", type=float, default=2.0,
                   help="Minimum experimental / control ratio of per-million gene cDNA (default 2)")
    p.add_argument("--pseudocount", type=float, default=1.0,
                   help="Pseudocount added to both group means, per million (default 1)")
    p.add_argument("-o", "--outdir", type=Path, default=REPO / "Centrosome", help="Output directory")
    return p.parse_args()


def run(cmd, stdout=None, stdin=None):
    # universal_newlines is used in place of text= for compatibility with Python 3.6.
    r = subprocess.run(cmd, stdout=stdout, stdin=stdin, stderr=subprocess.PIPE, universal_newlines=True)
    if r.returncode != 0:
        sys.exit("command failed: %s\n%s" % (" ".join(str(c) for c in cmd), r.stderr[:600]))
    return r


def require(tool):
    if shutil.which(tool) is None:
        sys.exit("%s not found on PATH" % tool)


def ucsc_chrom(name):
    chrom = name if name.startswith("chr") else "chr" + name
    return "chrM" if chrom == "chrMT" else chrom


def locate_peaks(raw_dir, workdir, prefixes):
    """Return {sample: path} for every <prefix>_R*_genome.*_Peaks.bed, extracting a tarball if needed."""
    tarballs = sorted(raw_dir.glob("*.tar.gz"))
    beds = sorted(raw_dir.glob("*_genome.*_Peaks.bed"))
    if not beds and len(tarballs) == 1:
        extract = workdir / "extracted"
        extract.mkdir(parents=True, exist_ok=True)
        with tarfile.open(str(tarballs[0])) as tf:
            tf.extractall(str(extract))
        beds = sorted(extract.rglob("*_genome.*_Peaks.bed"))
    found = {}
    for prefix in prefixes:
        hits = [b for b in beds if b.name.startswith(prefix + "_R")]
        if not hits:
            sys.exit("no %s_R*_genome.*_Peaks.bed found under %s" % (prefix, raw_dir))
        for b in hits:
            found[b.name.split("_genome.")[0]] = b
    return found


def gtf_to_genes(gtf, out):
    """Write the GTF's gene rows as BED6 (name = symbol, score = gene_id), primary contigs only."""
    awk = (
        'BEGIN{FS=OFS="\\t"} !/^#/ && $3=="gene" {'
        'sym=""; gid=""; '
        'if (match($9, /gene_name "[^"]+"/)) sym=substr($9, RSTART+11, RLENGTH-12); '
        'if (match($9, /gene_id "[^"]+"/)) gid=substr($9, RSTART+9, RLENGTH-10); '
        'if (sym=="") sym=gid; '
        'c=$1; if (c !~ /^chr/) c="chr"c; if (c=="chrMT") c="chrM"; '
        'print c, $4-1, $5, sym, gid, $7}'
    )
    raw = out.with_suffix(".unsorted")
    with open(str(raw), "w") as fh:
        run(["awk", awk, str(gtf)], stdout=fh)
    genes = {}
    with open(str(raw)) as fin, open(str(out), "w") as fout:
        for line in fin:
            c = line.rstrip("\n").split("\t")
            if c[0] not in PRIMARY:
                continue
            fout.write(line)
            genes[c[4]] = (c[3], c[0], int(c[1]), int(c[2]), c[5])
    raw.unlink()
    if not genes:
        sys.exit("no gene rows on primary contigs found in %s" % gtf)
    return genes


def normalise_peaks(src, dst, label):
    """Rewrite a peak file as sorted, chr-prefixed BED6 (column 4 = label), dropping chrM and scaffolds."""
    unsorted = dst.with_suffix(".unsorted")
    total = 0.0
    n = 0
    with open(str(src)) as fin, open(str(unsorted), "w") as fout:
        for line in fin:
            if not line.strip() or line.startswith(("#", "track", "browser")):
                continue
            c = line.rstrip("\n").split("\t")
            if len(c) < 6:
                sys.exit("%s: expected >=6 columns, got %d" % (src.name, len(c)))
            if c[0] in CHRM:
                continue
            chrom = ucsc_chrom(c[0])
            if chrom not in PRIMARY:
                continue
            start, end = int(c[1]), int(c[2])
            if end <= start:
                end = start + 1
            fout.write("\t".join([chrom, str(start), str(end), label, c[4], c[5]]) + "\n")
            total += float(c[4])
            n += 1
    with open(str(dst), "w") as fout:
        run(["sort", "-k1,1", "-k2,2n", str(unsorted)], stdout=fout)
    unsorted.unlink()
    return n, total


def gene_cdna(peaks, genes_bed):
    """cDNA per gene_id: each peak weighted by the fraction of its width inside the gene (same strand)."""
    r = run(["bedtools", "intersect", "-s", "-wo", "-a", str(peaks), "-b", str(genes_bed)], stdout=subprocess.PIPE)
    per_gene = defaultdict(float)
    for ln in r.stdout.splitlines():
        f = ln.split("\t")
        width = int(f[2]) - int(f[1])
        per_gene[f[10]] += float(f[4]) * min(int(f[-1]), width) / width
    return per_gene


def read_deseq(path, min_lfc, max_p):
    """Return ({symbol: (log2FoldChange, pvalue)} for all rows, [candidate symbols], n_rows)."""
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    names = [str(h).strip().lower() if h is not None else "" for h in header]
    try:
        i_lfc, i_p = names.index("log2foldchange"), names.index("pvalue")
    except ValueError:
        sys.exit("--deseq: expected columns log2FoldChange and pvalue, found %s" % header)
    stats = {}
    n = 0
    for r in rows:
        if r[0] is None:
            continue
        n += 1
        sym = r[0]
        if isinstance(sym, (datetime.datetime, datetime.date)):
            sym = tuple(DATE_SYMBOLS.get((sym.month, sym.day), ["%s" % sym.date()]))
        else:
            sym = (str(sym).strip(),)
        lfc, p = r[i_lfc], r[i_p]
        if lfc is None or p is None:
            continue
        stats[sym] = (float(lfc), float(p))
    candidates = [s for s, (lfc, p) in stats.items() if lfc > min_lfc and p <= max_p]
    wb.close()
    return stats, candidates, n


def match_symbols(candidates, genes):
    """Map each candidate (a tuple of alternative symbols) to the gene_ids carrying it."""
    exact = defaultdict(list)
    folded = defaultdict(list)
    for gid, (sym, _, _, _, _) in genes.items():
        exact[sym].append(gid)
        folded[sym.lower()].append(gid)
    matched, unmatched = {}, []
    for cand in candidates:
        gids = []
        for alt in cand:
            gids = exact.get(alt) or folded.get(alt.lower()) or []
            if gids:
                break
        if gids:
            matched[cand] = gids
        else:
            unmatched.append("|".join(cand))
    return matched, unmatched


def main():
    args = parse_args()
    require("bedtools")
    require("sort")
    require("awk")
    outdir = args.outdir
    workdir = outdir / "work"
    workdir.mkdir(parents=True, exist_ok=True)

    print("[1/6] reading gene annotation (%s)" % args.gtf)
    genes_bed = workdir / "genes.bed"
    genes = gtf_to_genes(args.gtf, genes_bed)
    print("      %d genes on primary contigs" % len(genes))

    print("[2/6] normalising peak files (%s)" % args.raw_dir)
    sources = locate_peaks(args.raw_dir, workdir, [args.experimental, args.control])
    samples = sorted(sources)
    exp_samples = [s for s in samples if s.startswith(args.experimental + "_R")]
    ctl_samples = [s for s in samples if s.startswith(args.control + "_R")]
    peak_paths, totals = {}, {}
    for s in samples:
        dst = workdir / (s + ".chr.bed")
        n, total = normalise_peaks(sources[s], dst, s.split("_")[-1])
        peak_paths[s], totals[s] = dst, total
        print("      %-16s %7d peaks  cDNA %10.0f" % (s, n, total))
        if total <= 0:
            sys.exit("%s has no peak cDNA on primary contigs" % s)

    print("[3/6] assigning peak cDNA to genes (per million per sample)")
    cpm = {}
    for s in samples:
        per_gene = gene_cdna(peak_paths[s], genes_bed)
        cpm[s] = {gid: v * 1e6 / totals[s] for gid, v in per_gene.items()}
    exp_peaks_per_gene = Counter()
    for s in exp_samples:
        r = run(["bedtools", "intersect", "-s", "-wa", "-wb", "-a", str(peak_paths[s]), "-b", str(genes_bed)],
                stdout=subprocess.PIPE)
        for ln in r.stdout.splitlines():
            exp_peaks_per_gene[ln.split("\t")[10]] += 1

    print("[4/6] reading differential-expression list (%s)" % args.deseq)
    stats, candidates, n_rows = read_deseq(args.deseq, args.min_lfc, args.max_p)
    matched, unmatched = match_symbols(candidates, genes)
    with open(str(workdir / "unmatched_symbols.txt"), "w") as fh:
        fh.write("\n".join(unmatched) + ("\n" if unmatched else ""))
    in_list = {}
    for cand, gids in matched.items():
        for gid in gids:
            in_list[gid] = stats[cand]
    print("      %d rows; %d pass log2FoldChange > %g and pvalue <= %g; %d matched to %d genes, %d unmatched"
          % (n_rows, len(candidates), args.min_lfc, args.max_p, len(matched), len(in_list), len(unmatched)))

    print("[5/6] computing experimental / control ratios")
    pc = args.pseudocount
    seen = set()
    for s in samples:
        seen.update(cpm[s])
    rows = []
    passing = set()
    for gid in sorted(seen, key=lambda g: (genes[g][1], genes[g][2])):
        vals = [cpm[s].get(gid, 0.0) for s in samples]
        mean_exp = sum(cpm[s].get(gid, 0.0) for s in exp_samples) / len(exp_samples)
        mean_ctl = sum(cpm[s].get(gid, 0.0) for s in ctl_samples) / len(ctl_samples)
        ratio = (mean_exp + pc) / (mean_ctl + pc)
        listed = gid in in_list
        ok = listed and ratio >= args.min_ratio
        if ok:
            passing.add(gid)
        lfc, p = in_list.get(gid, (float("nan"), float("nan")))
        rows.append([genes[gid][0], gid, int(listed), "%.4g" % lfc, "%.4g" % p]
                    + ["%.4f" % v for v in vals]
                    + ["%.4f" % mean_exp, "%.4f" % mean_ctl, "%.4f" % ratio, exp_peaks_per_gene.get(gid, 0), int(ok)])
    table = outdir / ("%s_gene_enrichment.tsv" % args.experimental)
    with open(str(table), "w") as fh:
        fh.write("\t".join(["symbol", "gene_id", "in_list", "log2FoldChange", "pvalue"] + samples
                           + ["mean_experimental", "mean_control", "ratio", "n_experimental_peaks", "pass"]) + "\n")
        for r in rows:
            fh.write("\t".join(str(x) for x in r) + "\n")
    listed_with_signal = sum(1 for gid in in_list if gid in seen)
    listed_with_exp = sum(1 for gid in in_list if exp_peaks_per_gene.get(gid, 0) > 0)
    all_cpm = sorted(v for gid in seen for s in exp_samples for v in [cpm[s].get(gid, 0.0)] if v > 0)
    median_cpm = all_cpm[len(all_cpm) // 2] if all_cpm else float("nan")
    print("      pseudocount %g per million = %.2f cDNA in the shallowest sample (%s); median non-zero gene value %.2f"
          % (pc, pc * min(totals.values()) / 1e6, min(totals, key=totals.get), median_cpm))
    print("      listed genes with any peak: %d; with an experimental peak: %d; passing ratio >= %g: %d"
          % (listed_with_signal, listed_with_exp, args.min_ratio, len(passing)))

    print("[6/6] writing anchors")
    passing_bed = workdir / "passing_genes.bed"
    with open(str(genes_bed)) as fin, open(str(passing_bed), "w") as fout:
        for line in fin:
            if line.split("\t")[4] in passing:
                fout.write(line)
    pooled = workdir / "experimental_pooled.bed"
    with open(str(pooled), "w") as fout:
        run(["sort", "-k1,1", "-k2,2n"] + [str(peak_paths[s]) for s in exp_samples], stdout=fout)
    merged = workdir / "experimental_merged.bed"
    with open(str(merged), "w") as fout:
        run(["bedtools", "merge", "-i", str(pooled), "-s", "-c", "4,5,6", "-o", "distinct,sum,distinct"], stdout=fout)
    anchors_all = workdir / "experimental_anchors.bed"
    n_merged = 0
    with open(str(merged)) as fin, open(str(anchors_all), "w") as fout:
        for line in fin:
            c = line.rstrip("\n").split("\t")
            n_merged += 1
            mid = (int(c[1]) + int(c[2])) // 2
            fout.write("\t".join([c[0], str(mid), str(mid + 1), c[3], c[4], c[5].split(",")[0]]) + "\n")
    r = run(["bedtools", "intersect", "-s", "-u", "-a", str(anchors_all), "-b", str(passing_bed)],
            stdout=subprocess.PIPE)
    support = Counter()
    tag = "min%gratio" % args.min_ratio
    out_bed = outdir / ("%s_%s_anchors.bed" % (args.experimental, tag))
    seen_keys = set()
    i = 0
    with open(str(out_bed), "w") as fout:
        for ln in r.stdout.splitlines():
            c = ln.split("\t")
            key = (c[0], c[1], c[5])
            if key in seen_keys:
                continue
            seen_keys.add(key)
            n = len(set(c[3].split(",")))
            support[n] += 1
            i += 1
            fout.write("\t".join([c[0], c[1], c[2], "%s_%dreps_%d" % (args.experimental, n, i), str(n), c[5]]) + "\n")
    print("      merged experimental regions: %d; anchors in retained genes: %d" % (n_merged, i))
    for n in sorted(support):
        print("        supported by %d/%d replicates: %d" % (n, len(exp_samples), support[n]))
    print("      -> %s\n      -> %s" % (out_bed, table))
    if i == 0:
        sys.exit("no anchors written; relax --min-ratio, --min-lfc or --max-p")


if __name__ == "__main__":
    main()
